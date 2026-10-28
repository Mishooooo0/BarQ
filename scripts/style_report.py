"""Build the branded, editable Word report and a self-contained print layout.

Run after draw_report_assets.py. Convert the generated HTML to PDF with Chromium.
Original character sprites and browser captures are embedded without alteration.
"""
from pathlib import Path
import base64
import json
import re
import subprocess
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
ASSETS = DOCS / 'report-assets'
SOURCE = DOCS / 'AI-FEATURE-PROJECT-REPORT.md'
POSES = ['hero', 'idea', 'think', 'receipt', 'think', 'places', 'idea', 'receipt', 'places', 'places', 'idea', 'celebrate', 'ok']


def data_url(path):
    mime = 'image/png' if path.suffix == '.png' else 'image/svg+xml'
    return f'data:{mime};base64,' + base64.b64encode(path.read_bytes()).decode()


def reference_document(path):
    d = Document()
    s = d.sections[0]
    s.page_width = Inches(8.27)
    s.page_height = Inches(11.69)
    s.left_margin = s.right_margin = Inches(.78)
    s.top_margin = Inches(.85)
    s.bottom_margin = Inches(.78)
    for name in ['Normal', 'Body Text', 'First Paragraph']:
        if name not in d.styles:
            continue
        style = d.styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string('090909')
        style.paragraph_format.line_spacing = 1.12
        style.paragraph_format.space_after = Pt(7)
    for name, size in [('Heading 1', 26), ('Heading 2', 17), ('Heading 3', 12)]:
        style = d.styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string('090909')
        style.paragraph_format.keep_with_next = True
    d.save(path)


def background(header):
    run = header.paragraphs[0].add_run()
    inline = run.add_picture(str(ASSETS / 'logo-watermark.png'), width=Inches(6.2), height=Inches(8))._inline
    anchor = OxmlElement('wp:anchor')
    for key, val in {'distT':'0', 'distB':'0', 'distL':'0', 'distR':'0', 'simplePos':'0', 'relativeHeight':'0', 'behindDoc':'1', 'locked':'0', 'layoutInCell':'1', 'allowOverlap':'1'}.items():
        anchor.set(key, val)
    simple = OxmlElement('wp:simplePos')
    simple.set('x','0'); simple.set('y','0'); anchor.append(simple)
    for axis, offset in [('H', .95), ('V', 2.0)]:
        pos=OxmlElement('wp:position'+axis);pos.set('relativeFrom','page')
        off=OxmlElement('wp:posOffset');off.text=str(int(Inches(offset)));pos.append(off);anchor.append(pos)
    anchor.append(deepcopy(inline.find(qn('wp:extent'))))
    effect=OxmlElement('wp:effectExtent')
    for k in ['l','t','r','b']:effect.set(k,'0')
    anchor.append(effect);anchor.append(OxmlElement('wp:wrapNone'))
    for tag in ['wp:docPr','wp:cNvGraphicFramePr','a:graphic']:
        element=inline.find(qn(tag))
        if element is not None:anchor.append(deepcopy(element))
    inline.getparent().replace(inline,anchor)


def build_word():
    ref=Path('/tmp/barq-themed-reference.docx')
    reference_document(ref)
    out=DOCS/'BarQ-AI-Feature-Report.docx'
    subprocess.run(['pandoc',str(SOURCE),'-f','markdown-implicit_figures','-t','docx','--standalone',f'--resource-path={DOCS}',f'--reference-doc={ref}','-o',str(out)],check=True)
    doc=Document(out)
    # Add the original hero sprite beside the cover mark.
    first_shape=doc.inline_shapes[0]
    ratio=first_shape.height/first_shape.width
    first_shape.width=Inches(4.7);first_shape.height=int(first_shape.width*ratio)
    cover_p=first_shape._inline.xpath('ancestor::w:p')[0]
    from docx.text.paragraph import Paragraph
    p=Paragraph(cover_p,doc._body)
    p.add_run('  ').add_picture(str(ASSETS/'character-hero.png'),width=Inches(1.08))
    for p in doc.paragraphs:
        if p.style.name=='Heading 2':
            match=re.match(r'(\d+)\.',p.text)
            if not match:continue
            number=int(match[1])
            p.paragraph_format.page_break_before=True
            props=p._p.get_or_add_pPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F9DF50');props.append(shade)
            p.add_run('  ').add_picture(str(ASSETS/f'character-{POSES[number-1]}.png'),width=Inches(.48))
    for table in doc.tables:
        screenshot_table=bool(table._tbl.xpath('.//w:drawing'))
        for i,row in enumerate(table.rows):
            props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
            if i==0:props.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                tc=cell._tc.get_or_add_tcPr()
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F9DF50' if i==0 else ('FFFFFF' if screenshot_table or i%2 else 'F5F5F7'));tc.append(shade)
                for para in cell.paragraphs:
                    if screenshot_table:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        run.font.size=Pt(9)
                        if i==0:run.bold=True
    for section in doc.sections:
        header=section.header
        p=header.paragraphs[0]
        p.text='BARQ  /  SMART SPENDING & HANGOUTS'
        for run in p.runs:run.font.size=Pt(8);run.font.bold=True
        background(header)
        p=section.footer.paragraphs[0];p.clear()
        p.add_run('AI FEATURE INNOVATION  •  STUDENT PROTOTYPE')
        p.add_run(' '*9)
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
        for run in p.runs:run.font.size=Pt(7);run.font.color.rgb=RGBColor.from_string('737682')
    doc.save(out)


def build_html():
    html=subprocess.check_output(['pandoc',str(SOURCE),'-f','markdown-implicit_figures','-t','html5'],text=True)
    # Embed the original files so the report can be opened offline.
    html=re.sub(r'src="(report-assets/[^\"]+)"',lambda m:'src="'+data_url(DOCS/m[1])+'"',html)
    chunks=re.split(r'(?=<h2\b)',html)
    cover=chunks[0]
    cover=cover.replace('<h1', '<div class="eyebrow">AI FEATURE INNOVATION / PROJECT REPORT</div><h1',1)
    # Cover gets a spacious character-and-mark composition.
    cover=re.sub(r'(<p><img[^>]+product-logo[^>]*>.*?</p>)',r'\1',cover)
    first_img=re.search(r'<p><img\b.*?</p>',cover,re.S)
    if first_img:
        image=first_img.group(0)
        replacement='<div class="cover-brand">'+image+f'<img class="cover-character" src="{data_url(ASSETS/"character-hero.png")}" alt="Barq character"></div>'
        cover=cover[:first_img.start()]+replacement+cover[first_img.end():]
    content='<section class="cover">'+cover+'</section>'
    for chunk in chunks[1:]:
        match=re.search(r'<h2[^>]*>(.*?)</h2>',chunk,re.S)
        title=re.sub('<[^>]+>','',match[1])
        num=int(re.match(r'(\d+)\.',title)[1])
        heading=match[0]
        top=f'<div class="chapter-top"><div><div class="eyebrow">BARQ / SECTION {num:02d}</div>{heading}</div><img class="chapter-character" src="{data_url(ASSETS/f"character-{POSES[num-1]}.png")}" alt="Barq character"></div>'
        chunk=chunk.replace(heading,top,1)
        chunk=re.sub(r'<table>(.*?)</table>',lambda m:'<table class="gallery">'+m[1]+'</table>' if '<img' in m[1] else m[0],chunk,flags=re.S)
        content+='<section class="chapter">'+chunk+'</section>'
    mark=json.loads((ASSETS/'brand-mark.json').read_text())
    watermark=f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="{mark["viewBox"]}"><path fill="#090909" d="{mark["path"]}"/></svg>'
    wm='data:image/svg+xml;base64,'+base64.b64encode(watermark.encode()).decode()
    css='''
    @page { size: A4; margin: 19mm 17mm 18mm; }
    * { box-sizing: border-box; }
    html { background: #f7f7f9; }
    body { margin: 0; color: #090909; font: 10.2pt/1.52 'DejaVu Sans', sans-serif; }
    .watermark { position: fixed; top: 47mm; left: 21mm; width: 130mm; height: 190mm; opacity: .035; z-index: -1; }
    .watermark img { width:100%; height:100%; }
    .edge {position:fixed;top:0;left:-8mm;width:2mm;height:255mm;background:#f9df50;border-radius:3mm;z-index:-1;}
    h1 {font-size:31pt;line-height:1.12;letter-spacing:-1.2px;margin:9mm 0 6mm;max-width:155mm;}
    h2 {font-size:21pt;line-height:1.2;letter-spacing:-.6px;margin:2mm 0 4mm;}
    h3 {font-size:11.5pt;line-height:1.35;margin:6mm 0 2mm;break-after:avoid;}
    p {margin:0 0 3.2mm;orphans:3;widows:3;}
    li {margin:0 0 2.2mm;} ul,ol {padding-left:5.5mm;margin-top:2mm;}
    strong {font-weight:700;} a {color:#53535a;text-decoration:underline;overflow-wrap:anywhere;}
    code {font:8.5pt 'DejaVu Sans Mono',monospace;overflow-wrap:anywhere;background:#ececf0;padding:.25mm .6mm;border-radius:1mm;}
    img {max-width:100%;object-fit:contain;}
    p:has(>img) {background:white;border:1px solid #e6e6e9;border-radius:4mm;padding:3mm;text-align:center;break-inside:avoid;}
    p>em:only-child {display:block;font-size:8pt;line-height:1.45;color:#66666f;}
    .eyebrow {font-size:8pt;letter-spacing:1.5px;font-weight:700;color:#6e6e77;}
    .cover {min-height:255mm;break-after:page;position:relative;padding-top:4mm;}
    .cover:before {content:'';display:block;width:24mm;height:3mm;background:#f9df50;border-radius:2mm;margin-bottom:8mm;}
    .cover h1 {max-width:150mm;}
    .cover-brand {display:flex;align-items:center;justify-content:space-between;margin:5mm 0 3mm;background:white;border-radius:5mm;padding:4mm;border:1px solid #e6e6e9;}
    .cover-brand p {border:0;padding:0;background:transparent;flex:1;margin:0;}
    .cover-brand p img {width:118mm !important;}
    .cover-character {width:35mm;height:42mm;object-fit:contain;}
    .cover>p:last-child {font-size:8.2pt;line-height:1.55;margin-top:6mm;padding-top:4mm;border-top:1px solid #dcdcde;}
    .chapter {break-before:page;}
    .chapter-top {display:flex;justify-content:space-between;align-items:center;border-bottom:3mm solid #f9df50;padding-bottom:2mm;margin-bottom:6mm;break-inside:avoid;break-after:avoid;}
    .chapter-top>div {flex:1;}
    .chapter-character {width:23mm;height:25mm;object-fit:contain;background:white;border-radius:3mm;padding:1mm;margin-left:4mm;}
    table {border-collapse:separate;border-spacing:0;width:100%;font-size:8.2pt;line-height:1.42;margin:4mm 0 6mm;border:1px solid #dedee4;border-radius:2mm;overflow:hidden;}
    th {background:#f9df50;color:#090909;font-weight:700;text-align:left;padding:2.5mm;}
    td {padding:2.5mm;vertical-align:top;border-top:1px solid #e4e4e9;overflow-wrap:anywhere;}
    tbody tr:nth-child(odd) {background:rgba(255,255,255,.88);} tbody tr:nth-child(even) {background:rgba(239,239,244,.9);}
    tr {break-inside:avoid;} thead {display:table-header-group;}
    .gallery {break-inside:avoid;table-layout:fixed;background:white;border-radius:4mm;}
    .gallery td {text-align:center;background:white;}
    .gallery img {width:66mm!important;height:auto;max-height:145mm;}
    .gallery th {text-align:center;font-size:10pt;}
    .gallery td p {margin:0;}
    .gallery tr:last-child td {font-size:8pt;line-height:1.45;color:#56565e;text-align:left;padding-bottom:3mm;}
    '''
    out='<!doctype html><html lang="en"><head><meta charset="utf-8"><title>BarQ — AI Feature Report</title><style>'+css+'</style></head><body><div class="watermark"><img alt="" src="'+wm+'"></div><div class="edge"></div>'+content+'</body></html>'
    (DOCS/'BarQ-AI-Feature-Report.html').write_text(out)


if __name__=='__main__':
    build_word()
    build_html()
    print('Built branded Word report and self-contained HTML print layout.')
